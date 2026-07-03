from __future__ import annotations

from docx import Document
from ebooklib import epub

from md_converter.cli import main


def _make_docx(path):
    document = Document()
    document.add_heading("CLI Test", level=1)
    document.add_paragraph("Contenu simple.")
    document.save(path)


def test_cli_converts_single_file(tmp_path):
    docx_path = tmp_path / "doc.docx"
    _make_docx(docx_path)
    output_path = tmp_path / "out.md"

    exit_code = main([str(docx_path), "-o", str(output_path)])

    assert exit_code == 0
    assert output_path.exists()
    assert "CLI Test" in output_path.read_text(encoding="utf-8")


def test_cli_batch_directory(tmp_path):
    src_dir = tmp_path / "src"
    src_dir.mkdir()
    _make_docx(src_dir / "a.docx")
    _make_docx(src_dir / "b.docx")

    out_dir = tmp_path / "out"

    exit_code = main([str(src_dir), "--output-dir", str(out_dir)])

    assert exit_code == 0
    assert (out_dir / "a.md").exists()
    assert (out_dir / "b.md").exists()


def test_cli_skips_existing_output_without_overwrite(tmp_path, capsys):
    docx_path = tmp_path / "doc.docx"
    _make_docx(docx_path)
    output_path = tmp_path / "out.md"
    output_path.write_text("déjà là", encoding="utf-8")

    exit_code = main([str(docx_path), "-o", str(output_path)])

    assert exit_code == 0
    assert output_path.read_text(encoding="utf-8") == "déjà là"


def test_cli_writes_epub_image_assets_next_to_markdown(tmp_path):
    book = epub.EpubBook()
    book.set_identifier("cli-image-test")
    book.set_title("Livre CLI")
    book.set_language("fr")

    image_bytes = b"\x89PNG\r\n\x1a\nfake-png-content"
    book.add_item(
        epub.EpubImage(
            uid="cover-img",
            file_name="images/illustration.png",
            media_type="image/png",
            content=image_bytes,
        )
    )

    chapter = epub.EpubHtml(title="Chapitre 1", file_name="chap_1.xhtml", lang="fr")
    chapter.content = '<h1>Chapitre 1</h1><img src="images/illustration.png" alt="illustration">'
    book.add_item(chapter)
    book.toc = (chapter,)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]

    epub_path = tmp_path / "livre.epub"
    epub.write_epub(str(epub_path), book)
    output_path = tmp_path / "out.md"

    exit_code = main([str(epub_path), "-o", str(output_path)])

    assert exit_code == 0
    asset_path = output_path.parent / "images" / "illustration.png"
    assert asset_path.exists()
    assert asset_path.read_bytes() == image_bytes
