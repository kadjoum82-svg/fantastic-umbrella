from __future__ import annotations

from docx import Document

from md_converter import convert


def _make_docx(path):
    document = Document()
    document.add_heading("Titre du document", level=1)

    paragraph = document.add_paragraph()
    run_bold = paragraph.add_run("Ce texte est en gras")
    run_bold.bold = True
    paragraph.add_run(" et ceci est normal.")

    document.add_paragraph("Premier élément", style="List Bullet")
    document.add_paragraph("Deuxième élément", style="List Bullet")

    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Colonne A"
    table.rows[0].cells[1].text = "Colonne B"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "2"

    document.save(path)


def test_convert_docx_headings_lists_and_tables(tmp_path):
    docx_path = tmp_path / "sample.docx"
    _make_docx(docx_path)

    result = convert(str(docx_path))

    assert "# Titre du document" in result.markdown
    assert "**Ce texte est en gras**" in result.markdown
    assert "- Premier élément" in result.markdown
    assert "- Deuxième élément" in result.markdown
    assert "| Colonne A | Colonne B |" in result.markdown
    assert "| 1 | 2 |" in result.markdown
