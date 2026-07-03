from __future__ import annotations

import base64
import io
import zipfile

from docx import Document
from ebooklib import epub

from md_converter.webapp.app import create_app


def _client():
    return create_app().test_client()


def test_index_serves_html_and_static_assets():
    client = _client()

    response = client.get("/")
    assert response.status_code == 200
    assert b"md-converter" in response.data

    for asset in ("style.css", "app.js", "manifest.webmanifest", "sw.js"):
        response = client.get(f"/static/{asset}")
        assert response.status_code == 200


def test_convert_route_rejects_empty_request():
    client = _client()

    response = client.post("/convert", data={}, content_type="multipart/form-data")

    assert response.status_code == 400
    assert "error" in response.get_json()


def test_convert_route_converts_uploaded_docx():
    document = Document()
    document.add_heading("Test Webapp", level=1)
    document.add_paragraph("Contenu de test pour l'app web.")
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    client = _client()
    response = client.post(
        "/convert",
        data={"file": (buffer, "test.docx")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    payload = response.get_json()
    assert payload["filename"] == "test.md"
    content = base64.b64decode(payload["content_b64"]).decode("utf-8")
    assert "Test Webapp" in content


def test_convert_route_bundles_epub_images_as_zip():
    book = epub.EpubBook()
    book.set_identifier("webapp-test")
    book.set_title("Livre WebApp")
    book.set_language("fr")

    image_bytes = b"\x89PNG\r\n\x1a\nfake"
    book.add_item(
        epub.EpubImage(uid="img", file_name="images/pic.png", media_type="image/png", content=image_bytes)
    )
    chapter = epub.EpubHtml(title="Chap 1", file_name="chap_1.xhtml", lang="fr")
    chapter.content = '<h1>Chap 1</h1><img src="images/pic.png">'
    book.add_item(chapter)
    book.toc = (chapter,)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]

    buffer = io.BytesIO()
    epub.write_epub(buffer, book)
    buffer.seek(0)

    client = _client()
    response = client.post(
        "/convert",
        data={"file": (buffer, "livre.epub")},
        content_type="multipart/form-data",
    )

    assert response.status_code == 200
    payload = response.get_json()
    assert payload["filename"] == "livre.zip"

    zip_bytes = base64.b64decode(payload["content_b64"])
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        names = zf.namelist()
        assert "livre.md" in names
        assert "images/pic.png" in names
        assert zf.read("images/pic.png") == image_bytes
